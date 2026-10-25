"""Bounded, offline document import. Parser dependencies are loaded on demand."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import stat
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from core.config import ImportConfig
from core.models import ExtractedRecord, ImportResult
from security.validators import UnsafePathError, safe_input_path
from security.parser_diagnostics import OMITTED_DIAGNOSTICS, private_parser_diagnostics
from security.office import OfficeValidationError, validate_office_archive


class DocumentImportError(ValueError):
    """An input could not be imported safely; no partial result is returned."""


FORMATS = {".csv", ".xlsx", ".pdf", ".docx", ".json", ".md", ".markdown", ".txt", ".html", ".htm"}


def _normalize(value: str) -> str:
    return unicodedata.normalize("NFC", value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _archive_check(data: bytes, config: ImportConfig) -> None:
    validate_office_archive(data, config)


def import_document(requested: str | Path, config: ImportConfig | None = None) -> ImportResult:
    config = config or ImportConfig()
    try:
        path = safe_input_path(requested)
        if path.suffix.lower() not in FORMATS:
            raise DocumentImportError("Unsupported format. Choose CSV, XLSX, PDF, DOCX, JSON, MD, TXT or HTML.")
        # Check the opened descriptor too, so a changing pathname cannot bypass size/type checks.
        descriptor = os.open(path, os.O_RDONLY | os.O_NONBLOCK | os.O_NOFOLLOW)
        with os.fdopen(descriptor, "rb") as handle:
            info = os.fstat(handle.fileno())
            if not stat.S_ISREG(info.st_mode):
                raise DocumentImportError("Choose a readable regular file, not a directory or device.")
            if info.st_size > config.max_file_bytes:
                raise DocumentImportError("File exceeds import.max_file_bytes; choose a smaller document.")
            # Reading the configured maximum preallocates a large buffer even
            # for tiny files. Use the checked size and one growth-detection byte.
            data = handle.read(info.st_size + 1)
            after = os.fstat(handle.fileno())
            if after.st_size > config.max_file_bytes or len(data) > config.max_file_bytes:
                raise DocumentImportError("File grew beyond the configured size limit.")
            if (len(data) != info.st_size or after.st_size != info.st_size
                    or after.st_mtime_ns != info.st_mtime_ns or after.st_ctime_ns != info.st_ctime_ns):
                raise DocumentImportError("Input changed while reading. Retry when the file is stable.")
        if len(data) > config.max_file_bytes:
            raise DocumentImportError("File grew beyond the configured size limit.")
        with private_parser_diagnostics() as diagnostics:
            result = _parse(path, data, config)
        if diagnostics.count:
            return ImportResult(result.source, result.records, [*result.warnings, OMITTED_DIAGNOSTICS], len(data))
        return ImportResult(result.source, result.records, result.warnings, len(data))
    except DocumentImportError:
        raise
    except OfficeValidationError as exc:
        raise DocumentImportError(str(exc)) from None
    except UnsafePathError as exc:
        raise DocumentImportError(str(exc)) from None
    except PermissionError:
        raise DocumentImportError('The selected file cannot be read. Check its Finder permissions or choose a readable copy.') from None
    except FileNotFoundError:
        raise DocumentImportError('The selected file is missing or moved. Select it again in Finder.') from None
    except zipfile.BadZipFile:
        raise DocumentImportError('The Office document archive is corrupt or unsupported. Obtain a fresh DOCX/XLSX export.') from None
    except UnicodeError:
        raise DocumentImportError('The text encoding cannot be read. Obtain a valid UTF-8 export.') from None
    except json.JSONDecodeError:
        raise DocumentImportError('The selected JSON is malformed. Obtain a valid UTF-8 JSON export.') from None
    except csv.Error:
        raise DocumentImportError('The selected CSV is malformed. Obtain a fresh CSV export with valid quoting and consistent columns.') from None
    except ImportError:
        raise DocumentImportError("Parser dependency missing. Run bash install.sh and retry.") from None
    except (OSError, ValueError, EOFError, KeyError, TypeError,
            IndexError, RecursionError, csv.Error, zipfile.BadZipFile):
        raise DocumentImportError(
            "Cannot read document. Check permissions, format and encoding; obtain a valid UTF-8 or original-format export."
        ) from None


def _parse(path: Path, data: bytes, config: ImportConfig) -> ImportResult:
    kind = path.suffix.lower().lstrip(".")
    source = str(path)
    records: list[ExtractedRecord] = []
    warnings: list[str] = []
    timestamp = datetime.now(timezone.utc)
    total_chars = 0

    def add(text: str, fields: dict[str, Any], location: dict[str, Any], title: str = "") -> None:
        nonlocal total_chars
        text = _normalize(text)
        if not text:
            return
        total_chars += len(text) + len(json.dumps(fields, ensure_ascii=False))
        if len(records) >= config.max_records or total_chars > config.max_text_chars:
            raise DocumentImportError("Extracted content exceeds import limits; split the document and retry.")
        identity = json.dumps([source, kind, text, fields, location], ensure_ascii=False,
                              sort_keys=True, allow_nan=False, separators=(",", ":"))
        records.append(ExtractedRecord(hashlib.sha256(identity.encode()).hexdigest(), source,
                       kind, _normalize(title) or path.name, text, fields, timestamp, location))

    def rows(values: Any, sheet: str | None = None) -> None:
        iterator = iter(values)
        header = next(iterator, None)
        if header is None:
            return
        headers: list[str] = []
        for index, value in enumerate(header, 1):
            label = _normalize(str(value)) if value is not None else ""
            label = label or f"column_{index}"
            original = label
            count = 2
            while label in headers:
                label = f"{original}_{count}"
                count += 1
            headers.append(label)
        for number, row in enumerate(iterator, 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(row) != len(headers):
                raise DocumentImportError(f"Row {number} has a different number of cells than the header.")
            fields = {key: _normalize(str(value)) if value is not None else ""
                      for key, value in zip(headers, row)}
            location: dict[str, Any] = {"row": number, "headers": headers}
            if sheet is not None:
                location["worksheet"] = sheet
            add("\n".join(f"{key}: {value}" for key, value in fields.items()), fields, location)

    if kind == "csv":
        # Decode incrementally instead of holding bytes + a complete Unicode
        # string + StringIO's character buffer alongside the extracted records.
        with io.TextIOWrapper(io.BytesIO(data), encoding="utf-8-sig", newline="") as text_stream:
            rows(csv.reader(text_stream, strict=True))
    elif kind == "xlsx":
        _archive_check(data, config)
        from openpyxl import load_workbook
        from lxml.etree import XMLSyntaxError
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    rows(sheet.iter_rows(values_only=True), sheet.title)
            finally:
                workbook.close()
        except (XMLSyntaxError, AttributeError):
            raise DocumentImportError("Malformed XLSX XML; obtain a valid spreadsheet export.") from None
        warnings.append("Spreadsheet formulas are not evaluated; only existing cached values are read.")
    elif kind == "docx":
        _archive_check(data, config)
        from docx import Document
        from lxml.etree import XMLSyntaxError
        try:
            document = Document(io.BytesIO(data))
            for number, paragraph in enumerate(document.paragraphs, 1):
                add(paragraph.text, {}, {"paragraph": number})
            for number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    fields = {f"column_{i}": cell.text for i, cell in enumerate(row.cells, 1)}
                    add("\n".join(fields.values()), fields, {"table": number, "row": row_number})
        except (XMLSyntaxError, AttributeError):
            raise DocumentImportError("Malformed DOCX XML; obtain a valid document export.") from None
    elif kind == "pdf":
        from importlib.metadata import version
        installed = version("pdfminer.six")
        if not installed.isdigit() or int(installed) < 20251230:
            raise DocumentImportError(
                "PDF parser is outdated or unrecognized; install the patched requirements before importing PDFs."
            )
        import pdfplumber
        from pdfplumber.page import Page
        from pdfminer.pdfpage import PDFPage
        from pdfplumber.utils.exceptions import PdfminerException
        from pdfminer.psparser import PSException
        from pdfminer.pdfdocument import PDFPasswordIncorrect
        try:
            with io.BytesIO(data) as stream:
                pdf = pdfplumber.open(stream)
                try:
                    if not pdf.doc.is_extractable:
                        raise DocumentImportError("PDF disallows text extraction; request an authorized export.")
                    pdf.doc.caching = False
                    doctop = 0
                    for number, raw_page in enumerate(PDFPage.create_pages(pdf.doc), 1):
                        if number > config.max_pdf_pages:
                            raise DocumentImportError("PDF exceeds import.max_pdf_pages; split it first.")
                        page = Page(pdf, raw_page, page_number=number, initial_doctop=doctop)
                        try:
                            text = page.extract_text() or ""
                            if len(text.strip()) < config.min_pdf_text_chars:
                                warnings.append(f"Page {number}: insufficient extractable text; may be scanned. No OCR performed.")
                            else:
                                add(text, {}, {"page": number})
                            doctop += page.height
                        finally:
                            page.close()
                finally:
                    # PDF.close() enumerates pdf.pages, even after an early limit
                    # failure. We own the stream and close each live page above;
                    # clear container caches without constructing every page again.
                    pdf.flush_cache()
        except PdfminerException as exc:
            if exc.args and isinstance(exc.args[0], PDFPasswordIncorrect):
                raise DocumentImportError('The selected PDF is password-protected. Request an authorized unlocked or text-based export.') from None
            raise DocumentImportError('The selected PDF is malformed, unreadable or encrypted. Obtain an authorized text-based export.') from None
        except PSException:
            raise DocumentImportError("Unreadable or encrypted PDF; request an authorized text-based export.") from None
    else:
        text = data.decode("utf-8-sig")
        if kind == "json":
            def reject_constant(value: str) -> None:
                raise DocumentImportError("JSON contains a nonstandard numeric constant.")
            parsed = json.loads(text, parse_constant=reject_constant)
            items = parsed if isinstance(parsed, list) else [parsed]
            for number, item in enumerate(items):
                fields = item if isinstance(item, dict) else {"value": item}
                add(json.dumps(fields, ensure_ascii=False, sort_keys=True, allow_nan=False), fields,
                    {"json_pointer": f"/{number}" if isinstance(parsed, list) else ""})
        elif kind in {"html", "htm"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(text, "html.parser")
            title = soup.title.get_text(" ", strip=True) if soup.title else ""
            for element in soup(["script", "style", "noscript", "template"]):
                element.decompose()
            add(soup.get_text("\n", strip=True), {}, {}, title)
        else:
            add(text, {}, {"line_start": 1})
    if not records:
        warnings.append("No usable records found. For scanned PDFs, request an authorized text-based export.")
    return ImportResult(source, records, warnings)
