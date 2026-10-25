"""Generated adversarial documents; no real documents, macros, or network use."""
from __future__ import annotations

import re
import socket
import stat
import zipfile
from unittest.mock import patch

import pytest

from core.config import ImportConfig
from handlers.local import DocumentImportError, import_document


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def deny(*args, **kwargs):
        raise AssertionError('document attempted a network operation')
    monkeypatch.setattr(socket, 'socket', deny)
    monkeypatch.setattr(socket, 'getaddrinfo', deny)


@pytest.fixture(params=['docx', 'xlsx'])
def office(request, tmp_path):
    path = tmp_path / ('fixture.' + request.param)
    if request.param == 'docx':
        from docx import Document
        doc = Document()
        doc.add_paragraph('Legitimate research fixture')
        doc.save(path)
    else:
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(['Name', 'Expertise'])
        book.active.append(['Test', 'Vision'])
        book.save(path)
        book.close()
    return path


def rewrite(path, replacements=None, additions=()):
    with zipfile.ZipFile(path) as original:
        members = [(info, original.read(info)) for info in original.infolist()]
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as changed:
        for info, data in members:
            if replacements and info.filename in replacements:
                data = replacements[info.filename](data)
            changed.writestr(info, data)
        for info, data in additions:
            changed.writestr(info, data)


def main_xml(path):
    return 'word/document.xml' if path.suffix == '.docx' else 'xl/worksheets/sheet1.xml'


@pytest.mark.parametrize('attack', ['external-file', 'external-network', 'external-dtd', 'expansion', 'utf16'])
def test_xml_entities_refused_before_third_party_parser(office, tmp_path, attack):
    sentinel = tmp_path / 'private.txt'
    sentinel.write_text('PRIVATE_ENTITY_SENTINEL')
    if attack == 'external-dtd':
        declaration = '<!DOCTYPE root SYSTEM "https://forbidden.test/private.dtd">'
    elif attack == 'external-network':
        declaration = '<!DOCTYPE root [<!ENTITY private SYSTEM "https://forbidden.test/private">]>'
    elif attack == 'expansion':
        declaration = '<!DOCTYPE root [<!ENTITY a "sensitive"><!ENTITY b "&a;&a;&a;">]>'
    else:
        declaration = f'<!DOCTYPE root [<!ENTITY private SYSTEM "{sentinel.as_uri()}">]>'
    def insert(data):
        text = re.sub(r'<\?xml[^>]*\?>', '', data.decode())
        reference = '&b;' if attack == 'expansion' else '&private;'
        if attack != 'external-dtd':
            text = text.replace('Legitimate research fixture', reference).replace('Vision', reference)
        text = declaration + text
        return text.encode('utf-16') if attack == 'utf16' else text.encode()
    rewrite(office, {main_xml(office): insert})
    parser = 'docx.Document' if office.suffix == '.docx' else 'openpyxl.load_workbook'
    with patch(parser, side_effect=AssertionError('unsafe XML reached document parser')):
        with pytest.raises(DocumentImportError, match='DTD|entities'):
            import_document(office)
    assert sentinel.read_text() == 'PRIVATE_ENTITY_SENTINEL'


@pytest.mark.parametrize('member', ['../escaped.txt', '/absolute.txt', 'word/../../escaped.txt',
                                    'word\\..\\escaped.txt', 'C:/escaped.txt', 'word//ambiguous.xml'])
def test_archive_traversal_refused_without_extraction(office, member):
    before = set(office.parent.iterdir())
    rewrite(office, additions=[(member, b'PRIVATE_ARCHIVE_CONTENT')])
    with pytest.raises(DocumentImportError, match='member path'):
        import_document(office)
    assert set(office.parent.iterdir()) == before


def test_symlink_archive_member_refused(office):
    info = zipfile.ZipInfo('linked.xml')
    info.create_system = 3
    info.external_attr = (stat.S_IFLNK | 0o777) << 16
    rewrite(office, additions=[(info, b'/etc/passwd')])
    with pytest.raises(DocumentImportError, match='symlink'):
        import_document(office)


def test_duplicate_archive_member_refused(office):
    with zipfile.ZipFile(office) as archive:
        data = archive.read(main_xml(office))
    with pytest.warns(UserWarning, match='Duplicate name'):
        rewrite(office, additions=[(main_xml(office), data)])
    with pytest.raises(DocumentImportError, match='Duplicate'):
        import_document(office)


def test_corrupt_xml_and_archive_are_actionable(office):
    rewrite(office, {main_xml(office): lambda data: b'<malformed PRIVATE_DOCUMENT_CONTENT'})
    with pytest.raises(DocumentImportError) as caught:
        import_document(office)
    assert 'PRIVATE_DOCUMENT_CONTENT' not in str(caught.value)
    office.write_bytes(office.read_bytes()[:40])
    with pytest.raises(DocumentImportError):
        import_document(office)


def test_archive_expansion_bounded_before_parsing(office):
    rewrite(office, additions=[('unused.bin', b'X' * 100_000)])
    with pytest.raises(DocumentImportError, match='Expanded'):
        import_document(office, ImportConfig(max_expanded_bytes=80_000))


def test_nested_xml_refused(office):
    namespace, tag = (('wordprocessingml', 'document') if office.suffix == '.docx'
                      else ('spreadsheetml', 'worksheet'))
    root = f'<{tag} xmlns="http://schemas.openxmlformats.org/{namespace}/2006/main">'.encode()
    xml = root + b'<a>' * 150 + b'</a>' * 150 + f'</{tag}>'.encode()
    rewrite(office, {main_xml(office): lambda data: xml})
    with pytest.raises(DocumentImportError, match='depth'):
        import_document(office)


@pytest.mark.parametrize('mutation', ['dimension', 'hidden-row', 'hidden-cell', 'shared-string'])
def test_malicious_xlsx_structures(tmp_path, mutation):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['Name', 'Expertise'])
    book.active.append(['Test', 'Vision'])
    path = tmp_path / 'malicious.xlsx'
    book.save(path)
    book.close()
    def mutate(data):
        if mutation == 'dimension':
            return data.replace(b'ref="A1:B2"', b'ref="A1:XFD1048576"')
        if mutation == 'hidden-row':
            return data.replace(b'<row r="2"', b'<row r="999999999"')
        if mutation == 'hidden-cell':
            return data.replace(b'r="B2"', b'r="XFD2"')
        return data.replace(b'<c r="A2" t="inlineStr"><is><t>Test</t></is></c>',
                            b'<c r="A2" t="s"><v>999999999</v></c>')
    rewrite(path, {'xl/worksheets/sheet1.xml': mutate})
    with pytest.raises(DocumentImportError):
        import_document(path)


@pytest.mark.parametrize('data', [b'%PDF-1.7\nPRIVATE_DOCUMENT_CONTENT',
    b'%PDF-1.7\n1 0 obj\n<< /Root 999 0 R >>\nendobj\n%%EOF',
    b'%PDF-1.7\nxref\nnot-a-number\ntrailer\n<<>>\n%%EOF'])
def test_malformed_pdfs_fail_without_sensitive_errors(tmp_path, data, caplog, capsys):
    path = tmp_path / 'broken.pdf'
    path.write_bytes(data)
    with pytest.raises(DocumentImportError) as caught:
        import_document(path)
    assert 'PRIVATE_DOCUMENT_CONTENT' not in str(caught.value)
    assert 'PRIVATE_DOCUMENT_CONTENT' not in caplog.text
    assert capsys.readouterr() == ('', '')


@pytest.mark.parametrize('suffix', ['pdf', 'docx', 'xlsx', 'txt'])
def test_oversized_inputs_rejected_before_parser(tmp_path, suffix):
    path = tmp_path / ('oversized.' + suffix)
    path.write_bytes(b'X' * 1025)
    with patch('handlers.local._parse', side_effect=AssertionError('size guard missed')):
        with pytest.raises(DocumentImportError, match='max_file_bytes'):
            import_document(path, ImportConfig(max_file_bytes=1024))


@pytest.mark.parametrize('path', ['../private.txt', '..\\private.txt'])
def test_input_traversal_is_actionable(tmp_path, monkeypatch, path):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(DocumentImportError):
        import_document(path)


def test_encrypted_archive_is_refused(office):
    data = bytearray(office.read_bytes())
    # ZIP flags occur at +6 in the local header, +8 in the central header.
    for signature, offset in ((b'PK\x03\x04', 6), (b'PK\x01\x02', 8)):
        index = data.index(signature) + offset
        data[index] |= 1
    office.write_bytes(data)
    with pytest.raises(DocumentImportError, match='Encrypted.*authorized'):
        import_document(office)


def test_xml_crc_corruption_is_actionable(office):
    data = bytearray(office.read_bytes())
    name = main_xml(office).encode()
    start = 0
    while True:
        header = data.find(b'PK\x01\x02', start)
        assert header >= 0
        name_length = int.from_bytes(data[header + 28:header + 30], 'little')
        if data[header + 46:header + 46 + name_length] == name:
            data[header + 16:header + 20] = b'\x00' * 4
            break
        start = header + 46 + name_length
    office.write_bytes(data)
    with pytest.raises(DocumentImportError) as caught:
        import_document(office)
    assert 'PRIVATE' not in str(caught.value)


def test_office_external_relationships_are_not_fetched(office, tmp_path):
    sentinel = tmp_path / 'private.txt'
    sentinel.write_text('PRIVATE_RELATIONSHIP_SENTINEL')
    relation = (f'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                f'TargetMode="External" Id="rIdPrivate" Target="{sentinel.as_uri()}"/>').encode()
    rewrite(office, {'_rels/.rels': lambda data: data.replace(b'</Relationships>', relation + b'</Relationships>')})
    result = import_document(office)
    assert result.records
    assert 'PRIVATE_RELATIONSHIP_SENTINEL' not in str(result.records)


def test_xlsx_column_limit_is_configurable(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['A', 'B'])
    book.active.append(['one', 'two'])
    path = tmp_path / 'columns.xlsx'
    book.save(path)
    book.close()
    with pytest.raises(DocumentImportError, match='dimensions'):
        import_document(path, ImportConfig(max_spreadsheet_columns=1))
    assert import_document(path, ImportConfig(max_spreadsheet_columns=2)).records[0].fields['B'] == 'two'


def test_well_formed_xml_with_wrong_document_structure_is_actionable(office):
    rewrite(office, {main_xml(office): lambda data: b'<unexpected/>'})
    with pytest.raises(DocumentImportError):
        import_document(office)
