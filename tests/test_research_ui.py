from dataclasses import replace
from pathlib import Path

import pytest

from core.config import AppConfig,CrawlConfig,OutputConfig,ImportConfig
from core.storage import RecordStore
from handlers.local import import_document
from research.models import ResearchProject
from research.storage import ProjectStore
from ui import menu,native,importing,research


def sequence(monkeypatch,name,values):
    answers=iter(values);monkeypatch.setattr(menu,name,lambda *a,**k:next(answers))


def test_project_creation_review_and_cancel(tmp_path,monkeypatch):
    api=ProjectStore(RecordStore(tmp_path/'output'))
    sequence(monkeypatch,'prompt_text',['Phishing WebAssembly','','phishing,WebAssembly','gaming'])
    sequence(monkeypatch,'prompt_choice',['Balanced'])
    monkeypatch.setattr(menu,'confirm',lambda _:True)
    project=research.create(api)
    assert project.keywords==('phishing','WebAssembly') and api.get(project.project_id).mode=='balanced'
    sequence(monkeypatch,'prompt_text',['Discarded','','',''])
    sequence(monkeypatch,'prompt_choice',[None])
    assert research.create(api) is None and len(api.list())==1


def test_project_document_import_association(tmp_path,monkeypatch,capsys):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    project=ResearchProject('Phishing',keywords=('phishing',));api.create(project)
    path=tmp_path/'source.txt';path.write_text('phishing research')
    monkeypatch.setattr(native,'select_document',lambda:str(path))
    monkeypatch.setattr(menu,'prompt_choice',lambda *args:None)
    importing.run(ImportConfig(),lambda:records,project=project)
    assert len(api.search(project.project_id))==1
    assert len(api.history(project.project_id))==1
    assert 'Project links: 1' in capsys.readouterr().out


def test_project_dashboard_cannot_bypass_failed_selfcheck(tmp_path,monkeypatch,capsys):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    p=ResearchProject('Phishing');api.create(p)
    sequence(monkeypatch,'prompt_choice',['Collect authorized websites','Back'])
    monkeypatch.setattr('ui.collection.run',lambda *a,**k:pytest.fail('collection despite failed self-check'))
    research.dashboard(api,p,AppConfig(CrawlConfig(),OutputConfig()),lambda:records,tmp_path/'config.toml',tmp_path,security_ok=False)
    assert 'disabled' in capsys.readouterr().out


def test_export_extension_prompt_and_safe_overwrite(tmp_path,monkeypatch):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    p=ResearchProject('Phishing');api.create(p)
    path=tmp_path/'source.txt';path.write_text('phishing research');api.ingest(p,import_document(path))
    monkeypatch.setattr(menu,'prompt_text',lambda _: 'findings')
    monkeypatch.setattr(menu,'prompt_choice',lambda *args:'json')
    monkeypatch.setattr(menu,'confirm',lambda _:False)
    research.export(api,p)
    before=(records.root/'findings.json').read_bytes()
    monkeypatch.setattr(menu,'confirm',lambda _:False)
    research.export(api,p)
    assert (records.root/'findings.json').read_bytes()==before


def test_readable_view_preserves_linebreaks_joiners_and_no_fabricated_summary(tmp_path,monkeypatch,capsys):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    p=ResearchProject('ශ්‍රී phishing',keywords=('phishing',));api.create(p)
    path=tmp_path/'source.txt';path.write_text('phishing\nශ්‍රී\nதமிழ்\x1b[31m\u202e')
    api.ingest(p,import_document(path))
    monkeypatch.setattr(menu,'confirm',lambda _:False)
    research.read(api.search(p.project_id)[0],p)
    output=capsys.readouterr().out
    assert 'No generated summary available' in output
    assert 'phishing\nශ්‍රී\nதமிழ்' in output
    assert '\x1b' not in output and '\u202e' not in output


def test_original_and_structured_content_quality_is_traceable(tmp_path):
    from processing.content import describe
    from core.text import tokenize
    path=tmp_path/'source.txt';path.write_text('ශ්‍රී\r\n\r\nதமிழ்\ufffd')
    result=import_document(path);detail=describe(result.records[0])
    assert detail['warnings']
    assert detail['original_sha256']
    assert tokenize('ශ්‍රී')==['ශ්‍රී']


def test_real_piped_project_import_read_and_markdown_export(tmp_path):
    import subprocess,sys,json
    import spiderdeck
    path=tmp_path/'source.txt';path.write_text('phishing research தமிழ் සිංහල')
    config=tmp_path/'config.toml'
    output=tmp_path/'data'
    config.write_text('[output]\ndir='+json.dumps(str(output)))
    answers=['0','1','Phishing research','','phishing','','1','y',
             '2',str(path),'4', '3','1','n','n','3',
             '5','project','n','1', '8','11','4','8']
    result=subprocess.run([sys.executable,'-B',str(Path(spiderdeck.__file__)),'--plain','--quiet','--config',str(config)],
        input='\n'.join(answers)+'\n',capture_output=True,text=True,timeout=20)
    assert result.returncode==0,result.stderr
    assert 'Research project created' in result.stdout
    assert 'SPIDERDECK RESEARCH DOCUMENT' in result.stdout
    assert 'Exported 1 research records' in result.stdout
    assert (output/'project.md').exists()
    records=RecordStore(output);assert records.count()==1
    api=ProjectStore(records)
    assert len(api.list())==1 and len(api.history(api.list()[0].project_id))==1


def test_docx_missing_core_properties_are_not_invented(tmp_path):
    import io,zipfile
    from xml.etree import ElementTree as ET
    from docx import Document
    path=tmp_path/'missing-metadata.docx'
    document=Document();document.add_paragraph('phishing research');document.save(path)
    original=path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source,zipfile.ZipFile(path,'w') as target:
        for info in source.infolist():
            if info.filename=='docProps/core.xml':continue
            data=source.read(info)
            if info.filename in {'_rels/.rels','[Content_Types].xml'}:
                root=ET.fromstring(data)
                for child in list(root):
                    if 'core-properties' in child.attrib.get('Type','') or child.attrib.get('PartName')=='/docProps/core.xml':
                        root.remove(child)
                data=ET.tostring(root)
            target.writestr(info,data)
    result=import_document(path)
    details=result.details[result.records[0].record_id]
    assert details['author'] is None and details['publication_date'] is None


def test_failed_project_import_is_recorded_without_sensitive_diagnostics(tmp_path,monkeypatch):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    p=ResearchProject('Phishing');api.create(p)
    path=tmp_path/'PRIVATE_filename.pdf';path.write_bytes(b'PRIVATE_document_not_pdf')
    monkeypatch.setattr(native,'select_document',lambda:str(path))
    importing.run(ImportConfig(),lambda:records,project=p)
    history=api.history(p.project_id)
    assert len(history)==1 and history[0]['status']=='collection_failed'
    assert 'PRIVATE' not in str(history)
    assert api.dashboard(api.get(p.project_id))['failures']==1


def test_docx_declared_title_and_headings_are_preserved(tmp_path):
    from docx import Document
    path=tmp_path/'title.docx'
    doc=Document();doc.core_properties.title='Declared research title';doc.add_heading('Research heading',level=1);doc.save(path)
    result=import_document(path)
    structure=result.details[result.records[0].record_id]
    assert structure['document_title']=='Declared research title'
    assert structure['headings']==[{'level':1,'text':'Research heading'}]


def test_project_selector_sanitizes_untrusted_titles(tmp_path,monkeypatch):
    records=RecordStore(tmp_path/'output');api=ProjectStore(records)
    api.create(ResearchProject('\x1b[31mPhishing\u202e'))
    choices=iter(['Open project',None,'Back'])
    def choose(message,labels):
        if message=='Choose project:':
            assert all('\x1b' not in label and '\u202e' not in label for label in labels)
        return next(choices)
    monkeypatch.setattr(menu,'prompt_choice',choose)
    research.run(AppConfig(CrawlConfig(),OutputConfig()),lambda:records,tmp_path/'config.toml',tmp_path)
