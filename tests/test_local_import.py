import io
import json
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest

from core.config import ImportConfig, load_config
from handlers.local import DocumentImportError, import_document
from security.validators import UnsafePathError, safe_output_path


def write(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding='utf-8')
    return path


def test_csv_structured_unicode_and_stable_ids(tmp_path):
    path = write(tmp_path, 'people.csv', 'Name,Expertise,Expertise\nනිපුන්,தமிழ்,Vision\n')
    first = import_document(path).records[0]
    second = import_document(path).records[0]
    assert first.fields == {'Name': 'නිපුන්', 'Expertise': 'தமிழ்', 'Expertise_2': 'Vision'}
    assert first.metadata['row'] == 2
    assert first.record_id == second.record_id
    assert len(first.record_id) == 64
    assert first.to_dict()['extracted_at'].endswith('Z')
    json.dumps(first.to_dict(), allow_nan=False)
    assert first.source == str(path.resolve())


@pytest.mark.parametrize('suffix', ['txt', 'md'])
def test_plain_text(tmp_path, suffix):
    record = import_document(write(tmp_path, f'notes.{suffix}', '  hello\nதமிழ்  ')).records[0]
    assert record.text == 'hello\nதமிழ்'


def test_json_preserves_nested_fields(tmp_path):
    records = import_document(write(tmp_path, 'data.json', '[{"name":"A","tags":["NLP"]},3]')).records
    assert records[0].fields['tags'] == ['NLP']
    assert records[1].fields == {'value': 3}
    assert records[1].metadata['json_pointer'] == '/1'


def test_html_is_offline_and_removes_scripts(tmp_path):
    path = write(tmp_path, 'page.html', '<title>Research</title><script>SECRET</script><style>HIDDEN</style><p>Vision</p><img src="https://example.com/a">')
    with patch('socket.socket', side_effect=AssertionError('network prohibited')):
        record = import_document(path).records[0]
    assert record.title == 'Research'
    assert 'SECRET' not in record.text
    assert 'HIDDEN' not in record.text
    assert 'Vision' in record.text


def test_xlsx(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Supervisors'
    sheet.append(['Name', 'Expertise'])
    sheet.append(['A', 'Computer vision'])
    path = tmp_path / 'people.xlsx'
    book.save(path)
    record = import_document(path).records[0]
    assert record.fields['Name'] == 'A'
    assert record.metadata['worksheet'] == 'Supervisors'
    assert record.metadata['row'] == 2


def test_docx(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph('Research in computer vision')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'A'
    path = tmp_path / 'people.docx'
    document.save(path)
    result = import_document(path)
    assert result.records[0].metadata['paragraph'] == 1
    assert result.records[1].fields['column_2'] == 'A'


def pdf_bytes(text):
    stream = f'BT /F1 12 Tf 50 750 Td ({text}) Tj ET'.encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data = b'%PDF-1.4\n'
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f'{index} 0 obj\n'.encode() + obj + b'\nendobj\n'
    start = len(data)
    data += b'xref\n0 6\n0000000000 65535 f \n'
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    return data + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()


def test_pdf_and_scanned_warning(tmp_path):
    path = tmp_path / 'people.pdf'
    path.write_bytes(pdf_bytes('Research in computer vision'))
    record = import_document(path).records[0]
    assert 'computer vision' in record.text
    assert record.metadata['page'] == 1
    path.write_bytes(pdf_bytes(''))
    result = import_document(path)
    assert not result.records
    assert any('No OCR' in warning for warning in result.warnings)


@pytest.mark.parametrize(('name', 'data'), [('bad.json', '{'), ('bad.json', 'NaN'), ('bad.csv', 'a,b\n1\n'), ('bad.xlsx', 'broken'), ('bad.docx', 'broken'), ('bad.pdf', 'broken'), ('bad.exe', 'x')])
def test_malformed(tmp_path, name, data):
    with pytest.raises(DocumentImportError):
        import_document(write(tmp_path, name, data))


def test_empty_and_encoding(tmp_path):
    assert not import_document(write(tmp_path, 'empty.txt', '')).records
    path = tmp_path / 'bad.txt'
    path.write_bytes(b'\xff')
    with pytest.raises(DocumentImportError):
        import_document(path)


def test_limits_and_missing_inputs(tmp_path):
    path = write(tmp_path, 'large.txt', 'abcdef')
    with pytest.raises(DocumentImportError, match='max_file_bytes'):
        import_document(path, ImportConfig(max_file_bytes=3))
    with pytest.raises(DocumentImportError, match='Extracted content'):
        import_document(path, ImportConfig(max_text_chars=3))
    path = write(tmp_path, 'rows.csv', 'a\n1\n2\n')
    with pytest.raises(DocumentImportError):
        import_document(path, ImportConfig(max_records=1))
    for path in [tmp_path, tmp_path / 'missing.txt', Path('/dev/null')]:
        with pytest.raises(DocumentImportError):
            import_document(path)


def test_archive_limits_and_macros(tmp_path):
    path = tmp_path / 'bad.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', 'a' * 100)
    with pytest.raises(DocumentImportError, match='Expanded'):
        import_document(path, ImportConfig(max_expanded_bytes=10))
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/vbaProject.bin', 'macro')
    with pytest.raises(DocumentImportError, match='Macro'):
        import_document(path)


def test_output_sibling_and_symlink_escape(tmp_path):
    base = tmp_path / 'output'
    base.mkdir()
    sibling = tmp_path / 'output-evil'
    sibling.mkdir()
    (base / 'link').symlink_to(sibling, target_is_directory=True)
    for requested in ['../output-evil/secret', sibling / 'secret', 'link/secret', '..\\escape']:
        with pytest.raises(UnsafePathError):
            safe_output_path(base, requested)
    assert not (sibling / 'secret').exists()


def test_config():
    assert load_config().imports.max_file_bytes > 0
    with pytest.raises(ValueError):
        ImportConfig(max_records=0)


def test_menu_import(tmp_path, monkeypatch, capsys):
    import argparse
    import spiderdeck
    from ui import menu
    from fx import effects
    from core.config import AppConfig, CrawlConfig, OutputConfig
    monkeypatch.setattr('core.config.load_config', lambda: AppConfig(
        CrawlConfig(), OutputConfig(str(tmp_path / 'output'))))
    path = write(tmp_path, 'notes.txt', 'Research')
    choices = iter(['2) Import a local document', '8) Quit'])
    monkeypatch.setattr(menu, 'main_menu', lambda: next(choices))
    monkeypatch.setattr(menu, 'prompt_path', lambda _: str(path))
    monkeypatch.setattr(menu, 'prompt_choice', lambda *args: 'Return to main menu')
    monkeypatch.setattr(spiderdeck, 'run_security_check', lambda: True)
    monkeypatch.setattr(effects, 'matrix_intro', lambda **_: None)
    monkeypatch.setattr(effects, 'glitch_title', lambda *_, **__: None)
    spiderdeck.run_menu(argparse.Namespace(no_fx=True, dry_run=False))
    assert 'Imported 1 records' in capsys.readouterr().out


def test_pdf_permission_enforced(tmp_path):
    from unittest.mock import MagicMock
    path = tmp_path / 'restricted.pdf'
    path.write_bytes(pdf_bytes('Do not extract'))
    document = MagicMock()
    document.doc.is_extractable = False
    with patch('pdfplumber.open') as opened:
        opened.return_value = document
        with pytest.raises(DocumentImportError, match='disallows text extraction'):
            import_document(path)
        document.pages.__iter__.assert_not_called()


def test_archive_member_limit(tmp_path):
    path = tmp_path / 'many.xlsx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('a', 'a')
        archive.writestr('b', 'b')
    with pytest.raises(DocumentImportError, match='too many members'):
        import_document(path, ImportConfig(max_archive_members=1))


def test_import_errors_do_not_reveal_document_contents(tmp_path):
    path = write(tmp_path, 'private.json', '{SECRET_PERSONAL_INFORMATION')
    with pytest.raises(DocumentImportError) as raised:
        import_document(path)
    assert 'SECRET' not in str(raised.value)


@pytest.mark.parametrize('installed', ['20231228', '20251107', 'unknown'])
def test_outdated_pdf_stack_refused(tmp_path, installed):
    path = tmp_path / 'notes.pdf'
    path.write_bytes(pdf_bytes('Research'))
    with patch('importlib.metadata.version', return_value=installed):
        with pytest.raises(DocumentImportError, match='parser is outdated'):
            import_document(path)
